# -*- coding: utf-8 -*-
"""Stage27.0 与 Stage26.9A 基线的正式内容一致性核验（只读）。"""
from __future__ import annotations

import importlib.util
import re
import sys
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
BASE = ROOT / 'outputs' / 'paper' / '课程设计论文_文本最终收口版_Stage26.9A.docx'
NEW = ROOT / 'outputs' / 'paper' / '课程设计论文_工程证据增强版_Stage27.0.docx'
MD = ROOT / 'docs' / 'paper' / '课程设计论文_工程证据增强稿_Stage27.0.md'


def load_builder():
    spec = importlib.util.spec_from_file_location(
        'b', str(ROOT / 'scripts' / '24a_build_stage24_docx.py'))
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    mod.FIGROOT = ROOT / 'outputs' / 'figures'
    mod.FIG_MAIN = []
    mod.FIG_APP = []
    return mod


def body_paragraphs(path):
    """正文段落（排除目录条目与空段）。"""
    doc = Document(str(path))
    out = []
    for p in doc.paragraphs:
        if p.style.name.lower().startswith('toc'):
            continue
        if p.text.strip():
            out.append(p.text)
    return out


def para_texts(path):
    doc = Document(str(path))
    return [p.text for p in doc.paragraphs]


def cell_texts(path):
    doc = Document(str(path))
    out = []
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                out.append(c.text)
    return out


def main() -> int:
    mod = load_builder()
    ib = mod.self_check(BASE)
    inw = mod.self_check(NEW)

    print('=' * 78)
    print('%-24s %-14s %-14s' % ('项目', 'Stage26.9A', 'Stage27.0'))
    for key in ('paragraphs', 'tables', 'omml', 'drawings', 'algo_tbl', 'formula_tbl'):
        print('%-24s %-14s %-14s' % (key, ib[key], inw[key]))
    print('%-24s %-14s %-14s' % ('fig_caps', len(ib['fig_caps']), len(inw['fig_caps'])))
    print('%-24s %-14s %-14s' % ('tbl_caps', len(ib['tbl_caps']), len(inw['tbl_caps'])))
    print('%-24s %-14s %-14s' % ('seq_fields', ib['seq_fields'], inw['seq_fields']))
    print('%-24s %-14s %-14s' % ('static_eq', len(ib['static_eq']), len(inw['static_eq'])))
    print('%-24s %-14s %-14s' % ('Heading1', ib['styles'].get('Heading 1', 0),
                                 inw['styles'].get('Heading 1', 0)))
    print('%-24s %-14s %-14s' % ('Reference', ib['styles'].get('Reference', 0),
                                 inw['styles'].get('Reference', 0)))
    print('-' * 78)
    drift = [(k, ib['anchors'].get(k, 0), inw['anchors'].get(k, 0))
             for k in sorted(set(ib['anchors']) | set(inw['anchors']))
             if ib['anchors'].get(k, 0) != inw['anchors'].get(k, 0)]
    print('正式锚点计数差异:', drift or '无（全部一致）')
    print('旧值扫描（基线）:', {k: v for k, v in ib['scan_old'].items() if v})
    print('旧值扫描（本轮）:', {k: v for k, v in inw['scan_old'].items() if v})
    print('禁用工程串（本轮）:', {k: v for k, v in inw['scan'].items() if v})
    print('-' * 78)

    # 表题清单一致性
    tb, tn = ib['tbl_caps'], inw['tbl_caps']
    print('表题清单完全一致:', tb == tn, '| 条数', len(tb), len(tn))
    if tb != tn:
        for a, b in zip(tb, tn):
            if a != b:
                print('   DIFF', a, '||', b)
    print('独立公式段落一致:', ib['eq_para_texts'] == inw['eq_para_texts'])
    print('公式制表位一致:', ib['eq_tabs'] == inw['eq_tabs'])

    # 图题：原有 21 张必须全部在（允许图号顺延）
    spec = importlib.util.spec_from_file_location(
        'b27', str(ROOT / 'scripts' / '37_stage27_0_paper_build.py'))
    b27 = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(b27)
    base_titles = [t for _, t, _, _ in b27.BASE_FIGS]
    new_titles = [re.sub(r'^图 \d+-\d+ ', '', c) for c in inw['fig_caps']]
    missing = [t for t in base_titles if t not in new_titles]
    print('原有正式图数:', len(base_titles), '| 新稿中缺失:', missing or '无')
    new_evidence = [t for _, t, _, _ in b27.EVIDENCE_FIGS]
    print('新增工程证据图数:', len(new_evidence), '| 缺失:',
          [t for t in new_evidence if t not in new_titles] or '无')
    print('图题总数:', len(new_titles))

    # 正文可见文本：排除目录条目后，原有段落必须逐一一致
    pb = body_paragraphs(BASE)
    pn = body_paragraphs(NEW)
    print('正文档落数 基线/本轮:', len(pb), len(pn))
    setn = set(pn)
    diff = [t for t in pb if t not in setn]
    print('基线段落不在新稿中的条目数:', len(diff))
    for t in diff[:20]:
        print('   MISSING:', t[:80])
    added = [t for t in pn if t not in set(pb)]
    print('新稿新增段落数:', len(added))
    return 0


if __name__ == '__main__':
    sys.exit(main())
